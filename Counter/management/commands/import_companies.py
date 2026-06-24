from django.core.management.base import BaseCommand
from django.db import transaction
import openpyxl
from Counter.models import Company, Province


class Command(BaseCommand):
    help = 'Importar empresas desde el archivo Empresas.xlsx'

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            type=str,
            default='Empresas.xlsx',
            help='Ruta del archivo xlsx a importar (default: Empresas.xlsx)'
        )

    @transaction.atomic
    def handle(self, *args, **options):
        file_path = options['file']
        self.stdout.write(f"Leyendo archivo: {file_path}")
        
        try:
            wb = openpyxl.load_workbook(file_path)
            ws = wb.active
            
            total_rows = ws.max_row - 1  # Excluir header
            created = 0
            updated = 0
            errors = 0
            
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 1):
                try:
                    provincia_name, ruc, empresa_name = row[0], str(int(row[1])), row[2]
                    
                    # Obtener o crear la provincia
                    provincia, _ = Province.objects.get_or_create(name=provincia_name)
                    
                    # Crear o actualizar la empresa
                    company, created_flag = Company.objects.get_or_create(
                        ruc=ruc,
                        defaults={'name': empresa_name, 'province': provincia}
                    )
                    
                    if created_flag:
                        created += 1
                    else:
                        # Actualizar si ya existe
                        company.name = empresa_name
                        company.province = provincia
                        company.save()
                        updated += 1
                    
                    # Mostrar progreso
                    if row_idx % 50 == 0:
                        self.stdout.write(f"Procesados {row_idx}/{total_rows}...")
                        
                except Exception as e:
                    errors += 1
                    self.stdout.write(
                        self.style.WARNING(f"Error en fila {row_idx + 1}: {str(e)}")
                    )
                    continue
            
            self.stdout.write(
                self.style.SUCCESS(
                    f"\n✅ Importación completada!\n"
                    f"  - Empresas creadas: {created}\n"
                    f"  - Empresas actualizadas: {updated}\n"
                    f"  - Errores: {errors}\n"
                    f"  - Total: {created + updated}"
                )
            )
            
        except FileNotFoundError:
            self.stdout.write(
                self.style.ERROR(f"❌ Archivo no encontrado: {file_path}")
            )
        except Exception as e:
            self.stdout.write(
                self.style.ERROR(f"❌ Error: {str(e)}")
            )
